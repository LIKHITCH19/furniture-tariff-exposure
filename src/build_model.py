"""Excel landed-cost model (output/landed_cost_model.xlsx).

Duty rates come from the Census data. The company inputs (volumes, sourcing
mix, freight) are placeholders. Results are all Excel formulas so the
workbook keeps working when the inputs change.
"""

from pathlib import Path

import duckdb
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data" / "processed" / "furniture_imports.csv"
OUT_FILE = ROOT / "output" / "landed_cost_model.xlsx"

# ---------------------------------------------------------------------------
# What the model covers
# ---------------------------------------------------------------------------

LINES = [  # hs6, label, annual purchases at China FOB prices (illustrative)
    ("940161", "Upholstered wood-frame seating", 12_000_000),
    ("940350", "Wooden bedroom furniture", 8_000_000),
    ("940360", "Other wooden furniture", 7_000_000),
    ("940340", "Kitchen cabinets and vanities", 5_000_000),
    ("940171", "Upholstered metal-frame seating", 4_000_000),
    ("940320", "Other metal furniture", 4_000_000),
]
ORIGINS = ["China", "Vietnam", "Malaysia", "Indonesia", "Cambodia", "Mexico"]

# Illustrative cost assumptions by origin - replace with real quotes.
FREIGHT = [0.07, 0.075, 0.075, 0.08, 0.08, 0.035]   # % of FOB
MPF_APPLIES = [1, 1, 1, 1, 1, 0]   # USMCA-originating goods are exempt
HMF_APPLIES = [1, 1, 1, 1, 1, 0]   # ocean freight only; Mexico moves by truck/rail
MPF_RATE, HMF_RATE = 0.003464, 0.00125

# Today's sourcing mix (illustrative), same origin order as ORIGINS
CURRENT_MIX = {
    "940161": [0.45, 0.40, 0.10, 0.05, 0, 0],
    "940350": [0.30, 0.55, 0.10, 0.05, 0, 0],
    "940360": [0.40, 0.45, 0.10, 0.05, 0, 0],
    "940340": [0.50, 0.40, 0.10, 0, 0, 0],
    "940171": [0.65, 0.35, 0, 0, 0, 0],
    "940320": [0.70, 0.25, 0, 0, 0, 0.05],
}

# Rates computed off less than this much trade are too thin to trust; those
# cells fall back to the origin's rate across all furniture lines.
MIN_VALUE = 2_000_000

# ---------------------------------------------------------------------------
# Styles (standard model conventions: blue = input, black = formula,
# green = link to another sheet, yellow fill = assumption to review)
# ---------------------------------------------------------------------------

FONT = "Arial"
F_BASE = Font(name=FONT, size=10)
F_INPUT = Font(name=FONT, size=10, color="0000FF")
F_LINK = Font(name=FONT, size=10, color="008000")
F_BOLD = Font(name=FONT, size=10, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True)
F_HEAD = Font(name=FONT, size=11, bold=True)
F_NOTE = Font(name=FONT, size=9, italic=True, color="595959")
F_WARN = Font(name=FONT, size=10, bold=True, color="C00000")
FILL_INPUT = PatternFill("solid", fgColor="FFFF00")
FILL_HEAD = PatternFill("solid", fgColor="D9E1F2")
FILL_TOTAL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)

USD = '$#,##0;($#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
PTS = '+0.0;-0.0;"-"'
IDX = '0.00'

O_COLS = [get_column_letter(4 + i) for i in range(len(ORIGINS))]   # D..I
FIRST_O, LAST_O = O_COLS[0], O_COLS[-1]


def put(ws, ref, value, font=F_BASE, fmt=None, fill=None, border=False, align=None):
    cell = ws[ref]
    cell.value = value
    cell.font = font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if border:
        cell.border = BOX
    if align:
        cell.alignment = align
    return cell


def header_row(ws, row, labels, start_col=2):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=start_col + i, value=label)
        c.font = F_BOLD
        c.fill = FILL_HEAD
        c.border = BOX
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ---------------------------------------------------------------------------
# Real duty rates from the Census data
# ---------------------------------------------------------------------------

def observed_rates():
    con = duckdb.connect()
    con.execute(f"""
        CREATE TABLE imports AS
        SELECT * FROM read_csv_auto('{DATA}', types={{'hs6': 'VARCHAR', 'country_code': 'VARCHAR'}})
    """)
    latest = con.execute("SELECT MAX(month) FROM imports").fetchone()[0]

    def rates(where):
        rows = con.execute(f"""
            SELECT country, hs6,
                   SUM(consumption_imports_usd) AS value,
                   SUM(calculated_duty_usd)     AS duty
            FROM imports
            WHERE {where}
            GROUP BY GROUPING SETS ((country, hs6), (country))
        """).fetchall()
        out = {}
        for country, hs6, value, duty in rows:
            out[(country, hs6 or "ALL")] = (value, duty / value if value else None)
        return out

    window = (f"month > DATE '{latest}' - INTERVAL 3 MONTH "
              f"AND month <= DATE '{latest}'")
    current = rates(window)
    base = rates("year = 2024")
    start = con.execute(f"SELECT MIN(month) FROM imports WHERE {window}").fetchone()[0]
    return current, base, start, latest


def pick(table, origin, hs6):
    """Rate for one origin x line, falling back to the origin's all-line rate."""
    value, rate = table.get((origin, hs6), (0, None))
    if value and value >= MIN_VALUE and rate is not None:
        return rate, value, False
    _, all_rate = table.get((origin, "ALL"), (0, None))
    return all_rate, value or 0, True


# ---------------------------------------------------------------------------
# Sheets
# ---------------------------------------------------------------------------

def sheet_readme(wb, start, latest):
    ws = wb.active
    ws.title = "Read me"
    widths(ws, {"A": 2, "B": 110})
    put(ws, "B2", "Furniture landed-cost scenario model", F_TITLE)
    lines = [
        ("What this answers", F_HEAD),
        ("If a mid-size US furniture importer moves volume out of China, what happens to its total "
         "landed cost once duty, freight and fees are counted - and how much more can a new "
         "factory charge before the move stops saving money?", F_BASE),
        ("", F_BASE),
        ("What is real and what is assumed", F_HEAD),
        ("REAL: duty rates. 'Duty rates' holds the effective rate US importers actually paid "
         f"(calculated duty / customs value) by origin and product line, "
         f"{start:%b %Y}-{latest:%b %Y}, from the US Census Bureau international trade API.", F_BASE),
        ("ASSUMED: the company. Purchase volumes, today's sourcing mix, freight costs and factory "
         "prices are illustrative inputs for a $40M-a-year importer. Replace them with real "
         "numbers before using this for a decision.", F_BASE),
        ("", F_BASE),
        ("How to use it", F_HEAD),
        ("1. Inputs - change purchase volumes, factory prices, freight, and override any duty rate.", F_BASE),
        ("2. Scenarios - edit today's mix. Scenarios 2-4 move today's China volume elsewhere and "
         "update on their own; type over them to define a different scenario.", F_BASE),
        ("3. Results - landed cost for each scenario and the change against today.", F_BASE),
        ("4. Break-even - the highest factory price premium over China at which each origin is "
         "still cheaper landed.", F_BASE),
        ("", F_BASE),
        ("Cell colours", F_HEAD),
        ("Blue text = a typed-in number (input or data). Black = formula. Green = pulled from "
         "another sheet. Yellow fill = an assumption worth reviewing.", F_BASE),
        ("", F_BASE),
        ("Things to know before trusting a number", F_HEAD),
        ("- Duty is applied to the FOB value. US customs value is normally the transaction value, "
         "which excludes international freight, so this matches how duty is assessed.", F_BASE),
        ("- Observed rates are averages across everything shipped under that product code. A "
         "single SKU can sit above or below them (exclusions, USMCA qualification, Section 232 "
         "coverage at the 10-digit level).", F_BASE),
        ("- Mexico's rate is a blend: USMCA-qualifying goods enter duty-free, others do not. "
         "Whether your goods qualify changes the answer a lot.", F_BASE),
        (f"- The rate window ({start:%b}-{latest:%b %Y}) straddles a policy change: Section 301 "
         "tariffs replaced the 10% Section 122 tariff on Jul 24, 2026. Use the overrides on "
         "'Inputs' to test the newer statutory rates.", F_BASE),
        (f"- Cells with under ${MIN_VALUE/1e6:.0f}M of trade in the window use the origin's rate "
         "across all furniture lines instead (marked with a comment on 'Duty rates').", F_BASE),
        ("- Merchandise processing fee (0.3464%) and harbor maintenance fee (0.125%) are applied "
         "as flat percentages; the per-entry MPF minimum and maximum are ignored.", F_BASE),
    ]
    for i, (text, font) in enumerate(lines, start=4):
        put(ws, f"B{i}", text, font, align=Alignment(wrap_text=True, vertical="top"))


def sheet_rates(wb, current, base, start, latest):
    ws = wb.create_sheet("Duty rates")
    widths(ws, {"A": 2, "B": 9, "C": 32, **{c: 13 for c in O_COLS}})
    put(ws, "B1", "Observed effective duty rates", F_TITLE)
    put(ws, "B2", "Source: US Census Bureau international trade API, imports for consumption. "
                  "Rate = calculated duty / customs value.", F_NOTE)

    blocks = [
        (4, f"Effective duty rate, {start:%b}-{latest:%b %Y} (used by the model)", current, PCT),
        (13, "Effective duty rate, calendar 2024 (before the 2025 tariffs)", base, PCT),
    ]
    fallbacks = []
    for top, title, table, fmt in blocks:
        put(ws, f"B{top}", title, F_HEAD)
        header_row(ws, top + 1, ["HS6", "Product line", *ORIGINS])
        for i, (hs6, label, _) in enumerate(LINES):
            r = top + 2 + i
            put(ws, f"B{r}", hs6, border=True)
            put(ws, f"C{r}", label, border=True)
            for col, origin in zip(O_COLS, ORIGINS):
                rate, value, fell_back = pick(table, origin, hs6)
                cell = put(ws, f"{col}{r}", rate, F_INPUT, fmt, border=True)
                if fell_back:
                    cell.comment = Comment(
                        f"Only ${value:,.0f} of {origin} trade in this line - "
                        f"using {origin}'s rate across all furniture lines instead.", "Likhit Vardhan Chiliveri")
                    if top == 4:
                        fallbacks.append(f"{origin} / {label}")

    put(ws, "B22", "Change since 2024 (percentage points)", F_HEAD)
    header_row(ws, 23, ["HS6", "Product line", *ORIGINS])
    for i, (hs6, label, _) in enumerate(LINES):
        r = 24 + i
        put(ws, f"B{r}", hs6, border=True)
        put(ws, f"C{r}", label, border=True)
        for col in O_COLS:
            put(ws, f"{col}{r}", f"=({col}{6 + i}-{col}{15 + i})*100", fmt=PTS, border=True)

    put(ws, "B31", f"Goods value behind each rate above, {start:%b}-{latest:%b %Y} ($)", F_HEAD)
    header_row(ws, 32, ["HS6", "Product line", *ORIGINS])
    for i, (hs6, label, _) in enumerate(LINES):
        r = 33 + i
        put(ws, f"B{r}", hs6, border=True)
        put(ws, f"C{r}", label, border=True)
        for col, origin in zip(O_COLS, ORIGINS):
            value = current.get((origin, hs6), (0, None))[0] or 0
            put(ws, f"{col}{r}", value, F_INPUT, USD, border=True)

    note = ("Rates fell back to the origin's all-line rate for: " + "; ".join(fallbacks)
            if fallbacks else "No cells needed a fallback rate.")
    put(ws, "B40", note, F_NOTE)
    ws.freeze_panes = "D4"


def sheet_inputs(wb):
    ws = wb.create_sheet("Inputs", 1)
    widths(ws, {"A": 2, "B": 9, "C": 44, **{c: 13 for c in O_COLS}, "D": 17})
    put(ws, "B1", "Inputs", F_TITLE)
    put(ws, "B2", "Yellow cells are assumptions for an illustrative importer - replace with real figures.", F_NOTE)

    # 1. purchase volumes
    put(ws, "B4", "1. What the company buys each year", F_HEAD)
    header_row(ws, 5, ["HS6", "Product line", "Purchases at China FOB prices ($)"])
    for i, (hs6, label, volume) in enumerate(LINES):
        r = 6 + i
        put(ws, f"B{r}", hs6, border=True)
        put(ws, f"C{r}", label, border=True)
        put(ws, f"D{r}", volume, F_INPUT, USD, FILL_INPUT, border=True)
    put(ws, "C12", "Total", F_BOLD, fill=FILL_TOTAL, border=True)
    put(ws, "D12", "=SUM(D6:D11)", F_BOLD, USD, FILL_TOTAL, border=True)
    ws["D5"].alignment = Alignment(wrap_text=True, horizontal="center")
    ws.row_dimensions[5].height = 42

    # 2. cost assumptions by origin
    put(ws, "B14", "2. Cost assumptions by origin", F_HEAD)
    header_row(ws, 15, ["", "Assumption", *ORIGINS])
    rows = [
        (16, "Factory price vs China (1.00 = same price)", [1.0] * len(ORIGINS), IDX),
        (17, "Freight and insurance (% of FOB)", FREIGHT, PCT),
        (18, "Merchandise processing fee applies (1 = yes)", MPF_APPLIES, "0"),
        (19, "Harbor maintenance fee applies (1 = yes)", HMF_APPLIES, "0"),
    ]
    for r, label, values, fmt in rows:
        put(ws, f"C{r}", label, border=True)
        for col, v in zip(O_COLS, values):
            put(ws, f"{col}{r}", v, F_INPUT, fmt, FILL_INPUT, border=True)
    ws["D16"].comment = Comment("China is the reference price. Enter supplier quotes as a ratio "
                                "to the China price for the same goods, e.g. 1.06 = 6% dearer.", "Likhit Vardhan Chiliveri")
    ws["I18"].comment = Comment("USMCA-originating goods are exempt from the MPF. Set to 1 if "
                                "your Mexican goods do not qualify.", "Likhit Vardhan Chiliveri")

    put(ws, "B21", "3. Fee rates", F_HEAD)
    put(ws, "C22", "Merchandise processing fee (% of value)", border=True)
    put(ws, "D22", MPF_RATE, F_INPUT, "0.0000%", border=True)
    put(ws, "C23", "Harbor maintenance fee (% of value)", border=True)
    put(ws, "D23", HMF_RATE, F_INPUT, "0.000%", border=True)
    put(ws, "C24", "Fee rate applied, by origin", F_BOLD, border=True)
    for col in O_COLS:
        put(ws, f"{col}24", f"=$D$22*{col}18+$D$23*{col}19", fmt=PCT2, border=True)
    ws["D22"].comment = Comment("Source: 19 CFR 24.23 (MPF ad valorem rate 0.3464%).", "Likhit Vardhan Chiliveri")
    ws["D23"].comment = Comment("Source: 19 CFR 24.24 (HMF 0.125% of value, ocean cargo).", "Likhit Vardhan Chiliveri")

    # 4. overrides
    put(ws, "B26", "4. Duty rate overrides (leave blank to use the observed rate)", F_HEAD)
    header_row(ws, 27, ["HS6", "Product line", *ORIGINS])
    for i, (hs6, label, _) in enumerate(LINES):
        r = 28 + i
        put(ws, f"B{r}", hs6, border=True)
        put(ws, f"C{r}", label, border=True)
        for col in O_COLS:
            put(ws, f"{col}{r}", None, F_INPUT, PCT, FILL_INPUT, border=True)

    # 5. rate actually used
    put(ws, "B35", "5. Duty rate used in the model", F_HEAD)
    header_row(ws, 36, ["HS6", "Product line", *ORIGINS])
    for i, (hs6, label, _) in enumerate(LINES):
        r = 37 + i
        put(ws, f"B{r}", hs6, border=True)
        put(ws, f"C{r}", label, border=True)
        for col in O_COLS:
            formula = f"=IF({col}{28 + i}=\"\",'Duty rates'!{col}{6 + i},{col}{28 + i})"
            put(ws, f"{col}{r}", formula, F_LINK, PCT, border=True)
    ws.freeze_panes = "D4"


SCENARIOS = [
    ("Today's sourcing", None),
    ("Move China volume to Vietnam", {"Vietnam": 1.0}),
    ("Spread China volume: Vietnam / Cambodia / Indonesia", {"Vietnam": 0.4, "Cambodia": 0.3, "Indonesia": 0.3}),
    ("Nearshore China volume to Mexico", {"Mexico": 1.0}),
]
SCEN_TOP = [4 + 11 * s for s in range(len(SCENARIOS))]   # first row of each block


def sheet_scenarios(wb):
    ws = wb.create_sheet("Scenarios", 2)
    widths(ws, {"A": 2, "B": 9, "C": 34, **{c: 12 for c in O_COLS}, "J": 10, "K": 24})
    put(ws, "B1", "Sourcing scenarios", F_TITLE)
    put(ws, "B2", "Share of each product line bought from each origin. Scenarios 2-4 move today's "
                  "China share elsewhere - type over them to build your own.", F_NOTE)

    china = O_COLS[ORIGINS.index("China")]
    for s, (name, moves) in enumerate(SCENARIOS):
        top = SCEN_TOP[s]
        put(ws, f"B{top}", f"Scenario {s + 1}", F_HEAD)
        put(ws, f"C{top}", name, F_INPUT if s == 0 else F_BOLD, fill=FILL_INPUT if s == 0 else None)
        header_row(ws, top + 1, ["HS6", "Product line", *ORIGINS, "Total", "Check"])
        for i, (hs6, label, _) in enumerate(LINES):
            r = top + 2 + i
            base_r = SCEN_TOP[0] + 2 + i
            put(ws, f"B{r}", hs6, border=True)
            put(ws, f"C{r}", label, border=True)
            for j, (col, origin) in enumerate(zip(O_COLS, ORIGINS)):
                if s == 0:
                    put(ws, f"{col}{r}", CURRENT_MIX[hs6][j], F_INPUT, PCT, FILL_INPUT, border=True)
                elif origin == "China":
                    put(ws, f"{col}{r}", 0, F_INPUT, PCT, border=True)
                else:
                    share = moves.get(origin, 0)
                    f = f"={col}{base_r}" + (f"+{share}*${china}{base_r}" if share else "")
                    put(ws, f"{col}{r}", f, fmt=PCT, border=True)
            put(ws, f"J{r}", f"=SUM({FIRST_O}{r}:{LAST_O}{r})", F_BOLD, PCT, FILL_TOTAL, border=True)
            put(ws, f"K{r}", f'=IF(ABS(J{r}-1)>0.0001,"Mix must add to 100%","OK")', F_WARN, border=True)
    ws.freeze_panes = "D4"


RES_TOP = [14 + 10 * s for s in range(len(SCENARIOS))]


def sheet_results(wb):
    ws = wb.create_sheet("Results", 3)
    widths(ws, {"A": 2, "B": 9, "C": 46, "D": 15, "E": 15, "F": 15, "G": 13, "H": 16, "I": 13})
    put(ws, "B1", "Landed cost by scenario", F_TITLE)
    put(ws, "B2", "Landed cost = factory price + freight + duty + customs fees. Annual, US$.", F_NOTE)

    # summary
    put(ws, "B4", "Summary", F_HEAD)
    header_row(ws, 5, ["#", "Scenario", "Landed cost ($)", "vs today ($)", "vs today (%)",
                       "Duty paid ($)", "Effective duty rate", "FOB cost ($)"])
    for s in range(len(SCENARIOS)):
        r = 6 + s
        tot = RES_TOP[s] + 8
        put(ws, f"B{r}", s + 1, border=True)
        put(ws, f"C{r}", f"=Scenarios!C{SCEN_TOP[s]}", F_LINK, border=True)
        put(ws, f"D{r}", f"=H{tot}", fmt=USD, border=True)
        put(ws, f"E{r}", f"=D{r}-$D$6", fmt=USD, border=True)
        put(ws, f"F{r}", f"=IF($D$6=0,0,D{r}/$D$6-1)", fmt=PCT, border=True)
        put(ws, f"G{r}", f"=F{tot}", fmt=USD, border=True)
        put(ws, f"H{r}", f"=IF(D{tot}=0,0,F{tot}/D{tot})", fmt=PCT, border=True)
        put(ws, f"I{r}", f"=D{tot}", fmt=USD, border=True)
    put(ws, "C11", "Negative 'vs today' = the scenario is cheaper than today's sourcing.", F_NOTE)

    for s, (name, _) in enumerate(SCENARIOS):
        top = RES_TOP[s]
        put(ws, f"B{top}", f"Scenario {s + 1}", F_HEAD)
        put(ws, f"C{top}", f"=Scenarios!C{SCEN_TOP[s]}", Font(name=FONT, size=11, bold=True, color="008000"))
        header_row(ws, top + 1, ["HS6", "Product line", "FOB cost ($)", "Freight ($)", "Duty ($)",
                                 "Fees ($)", "Landed cost ($)", "Duty rate"])
        for i, (hs6, label, _) in enumerate(LINES):
            r = top + 2 + i
            mix = f"Scenarios!$D${SCEN_TOP[s] + 2 + i}:$I${SCEN_TOP[s] + 2 + i}"
            vol = f"Inputs!$D${6 + i}"
            price = "Inputs!$D$16:$I$16"
            put(ws, f"B{r}", hs6, border=True)
            put(ws, f"C{r}", label, border=True)
            put(ws, f"D{r}", f"={vol}*SUMPRODUCT({mix},{price})", fmt=USD, border=True)
            put(ws, f"E{r}", f"={vol}*SUMPRODUCT({mix},{price},Inputs!$D$17:$I$17)", fmt=USD, border=True)
            put(ws, f"F{r}", f"={vol}*SUMPRODUCT({mix},{price},Inputs!$D${37 + i}:$I${37 + i})",
                fmt=USD, border=True)
            put(ws, f"G{r}", f"={vol}*SUMPRODUCT({mix},{price},Inputs!$D$24:$I$24)", fmt=USD, border=True)
            put(ws, f"H{r}", f"=SUM(D{r}:G{r})", F_BOLD, USD, border=True)
            put(ws, f"I{r}", f"=IF(D{r}=0,0,F{r}/D{r})", fmt=PCT, border=True)
        t = top + 8
        put(ws, f"C{t}", "Total", F_BOLD, fill=FILL_TOTAL, border=True)
        for col in "DEFGH":
            put(ws, f"{col}{t}", f"=SUM({col}{top + 2}:{col}{top + 7})", F_BOLD, USD, FILL_TOTAL, border=True)
        put(ws, f"I{t}", f"=IF(D{t}=0,0,F{t}/D{t})", F_BOLD, PCT, FILL_TOTAL, border=True)
    ws.freeze_panes = "D4"


def sheet_breakeven(wb):
    ws = wb.create_sheet("Break-even", 4)
    alts = ORIGINS[1:]
    widths(ws, {"A": 2, "B": 9, "C": 34, **{get_column_letter(4 + i): 13 for i in range(len(alts))}})
    put(ws, "B1", "How much more can a new factory charge?", F_TITLE)
    put(ws, "B2", "The highest price premium over China at which each origin is still cheaper "
                  "landed (duty, freight and fees included). Negative = more expensive than China "
                  "even at the same factory price.", F_NOTE)
    header_row(ws, 4, ["HS6", "Product line", *alts])
    cn = "$D"   # China is the first origin column on Inputs
    for i, (hs6, label, _) in enumerate(LINES):
        r = 5 + i
        ir = 37 + i   # duty-rate-used row on Inputs
        put(ws, f"B{r}", hs6, border=True)
        put(ws, f"C{r}", label, border=True)
        china_cost = f"(1+Inputs!{cn}$17+Inputs!{cn}${ir}+Inputs!{cn}$24)"
        for j, origin in enumerate(alts):
            src = O_COLS[j + 1]     # same origin's column on Inputs
            dst = get_column_letter(4 + j)
            alt_cost = f"(1+Inputs!{src}$17+Inputs!{src}${ir}+Inputs!{src}$24)"
            put(ws, f"{dst}{r}", f"={china_cost}/{alt_cost}-1", fmt=PCT, border=True)
    put(ws, "C12", "Example: 20% in the Vietnam column means a Vietnamese factory can quote up "
                   "to 20% above the China price for the same goods and still come out cheaper "
                   "once the goods reach the US.", F_NOTE)
    ws["C12"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("C12:H14")


def main():
    current, base, start, latest = observed_rates()
    wb = Workbook()
    sheet_readme(wb, start, latest)
    sheet_inputs(wb)
    sheet_scenarios(wb)
    sheet_results(wb)
    sheet_breakeven(wb)
    sheet_rates(wb, current, base, start, latest)
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.properties.creator = "Likhit Vardhan Chiliveri"
    wb.properties.title = "Furniture landed-cost scenario model"

    # print each sheet landscape, one page wide
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.sheet_view.showGridLines = False

    OUT_FILE.parent.mkdir(exist_ok=True)
    wb.save(OUT_FILE)
    print(f"wrote {OUT_FILE.relative_to(ROOT)}  (rates from {start:%b %Y} to {latest:%b %Y})")


if __name__ == "__main__":
    main()
